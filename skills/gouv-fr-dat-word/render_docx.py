#!/usr/bin/env python3
"""Rend un document d'homologation MirAI au format ministériel `.docx`.

Construit le `.docx` **sur le modèle** `MODELE_DAT_MirAI.docx` (hérite page de garde,
styles, sommaire/TOC) à partir d'un contenu **Markdown fidèle à la structure du modèle**
(tel que produit par la skill `dat-generation`).

Opérations :
  - conserve la page de garde et le champ Sommaire (TOC) du modèle ;
  - supprime la page « Comment utiliser ce modèle » et tous les encadrés d'instruction (« ✎ ») ;
  - supprime le squelette de corps du modèle et le remplace par le contenu fourni ;
  - applique les styles de titre RÉELS du modèle (styleId Heading1/2/3 — l'accès par nom peut échouer) ;
  - saut de page avant chaque section de niveau 1 ; `keep_with_next` sur les titres ;
  - tableaux Markdown → tableaux Word (en-tête répété `tblHeader`, lignes non sécables `cantSplit`) ;
  - blocs de code ``` → police à chasse fixe (rend l'ASCII des schémas) ;
  - marque le TOC « dirty » + `settings/updateFields=true` (sinon sommaire vide ;
    repli utilisateur : Ctrl+A puis F9).

Mapping des niveaux Markdown → modèle : `##` → Heading1, `###` → Heading2, `####` → Heading3.
Le titre `#` et le bloc de méta en tête (avant le premier `##`) sont ignorés (déjà sur la garde).

Usage :
    python render_docx.py --input DAT.md --output private/DAT_x_v0.1.docx \
        [--model /chemin/MODELE_DAT_MirAI.docx] [--service "Nom"] [--version v0.1] [--date "21 juin 2026"]

Dépendance : python-docx (`pip install python-docx`).
"""
from __future__ import annotations

import argparse
import json
import os
import re
import shutil
import subprocess
import sys
import tempfile
from types import SimpleNamespace

try:
    from docx import Document
    from docx.oxml.ns import qn, nsdecls
    from docx.oxml import OxmlElement, parse_xml
    from docx.shared import Inches, Pt, Cm
    from docx.enum.text import WD_ALIGN_PARAGRAPH
except ImportError:
    sys.exit("python-docx requis : pip install python-docx")

NS_W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
LEVEL_STYLE = {2: "Heading1", 3: "Heading2", 4: "Heading3"}  # nb de '#' -> styleId modèle

# ── Mise en page (réglable) ───────────────────────────────────────────────────
# Le modèle est dense (corps 11 pt, H1 16 / H2 13 / H3 12, sans interligne custom).
# Ces valeurs réduisent les polices ET aèrent (interligne + espace après paragraphe).
BODY_PT = 10.0          # corps de texte
H1_PT, H2_PT, H3_PT = 14.0, 12.0, 11.0  # titres
CODE_PT = 8.5           # blocs de code à chasse fixe (ASCII résiduel)
TABLE_PT = 9.0          # cellules de tableau
VERSION_TABLE_PT = 8.0  # table « Suivi des mises à jour » (page de garde) — compacte
LINE_SPACING = 1.3      # interligne du corps (multiple)
PARA_AFTER_PT = 8.0     # espace après paragraphe du corps (aération / séparation)
HEADING_SPACING = {     # (avant, après) en points, par style de titre
    "Heading1": (12.0, 4.0),
    "Heading2": (10.0, 4.0),
    "Heading3": (8.0, 2.0),
}
MERMAID_MAX_H_IN = 7.0  # hauteur image max (pouces) avant downscale

# ── Listes (lisibilité) ───────────────────────────────────────────────────────
BULLET = "-  "          # puce de liste : trait d'union (choix ministériel)
LIST_INDENT_CM = 0.6    # retrait gauche des items
LIST_HANGING_CM = 0.4   # retrait pendant (lignes suivantes alignées sous le texte)
LIST_AFTER_PT = 3.0     # petit espace entre items

# ── Typographie française (Imprimerie nationale) ──────────────────────────────
NBSP = " "            # espace insécable
NNBSP = " "           # espace fine insécable
THIN_NBSP_FALLBACK = False  # True : replie toutes les fines U+202F sur U+00A0 (police « tofu »)
# Unités fréquentes : nombre + unité → insécable (si l'espace existe). Conservateur.
TYPO_UNITS = ["pt", "px", "cm", "mm", "km", "kg", "Mo", "Go", "To", "Ko", "ko",
              "min", "h", "j", "jours", "jour", "s", "m"]

# ── Logo de la page de garde ──────────────────────────────────────────────────
LOGO_DEFAULT = os.path.join(os.path.dirname(os.path.abspath(__file__)),
                            "assets", "logo_minint_2020.png")
LOGO_WIDTH_CM = 4.0

# Emplacements usuels d'un Chrome/Chromium pour le rendu mermaid (aucun service externe).
CHROME_CANDIDATES = [
    "/Applications/Google Chrome.app/Contents/MacOS/Google Chrome",
    "/Applications/Chromium.app/Contents/MacOS/Chromium",
    "/usr/bin/google-chrome",
    "/usr/bin/chromium",
    "/usr/bin/chromium-browser",
]


# ── Helpers XML sur le corps du modèle ────────────────────────────────────────

def p_text(el) -> str:
    return "".join(t.text or "" for t in el.findall(".//" + qn("w:t")))


def has_toc(el) -> bool:
    """True si l'élément (w:p ou w:sdt) contient un champ TOC dans ses descendants."""
    return any("TOC" in (it.text or "") for it in el.iter(qn("w:instrText")))


def strip_template_body(doc) -> None:
    """Retire la page how-to et le squelette ; conserve garde + Sommaire + TOC + sectPr.

    Robuste au TOC encapsulé dans un w:sdt (Word moderne) : les repères sont cherchés
    parmi les enfants DIRECTS du corps via leurs descendants textuels.
    """
    body = doc.element.body
    children = list(body)

    def find(pred):
        # cherche un enfant direct dont un descendant w:p correspond au prédicat
        for el in children:
            for p in ([el] if el.tag == qn("w:p") else el.iter(qn("w:p"))):
                if pred(p_text(p).strip()):
                    return el
        return None

    howto = find(lambda t: t.startswith("Comment utiliser ce mod"))
    sommaire = find(lambda t: t == "Sommaire")
    toc = next((el for el in children if has_toc(el)), None)
    sectpr = body.find(qn("w:sectPr"))

    if howto is None or sommaire is None or toc is None:
        raise SystemExit("Modèle inattendu : repères (how-to / Sommaire / TOC) introuvables.")

    i_howto, i_somm, i_toc = children.index(howto), children.index(sommaire), children.index(toc)

    # 1) supprime la page how-to (de 'Comment utiliser' jusqu'avant 'Sommaire')
    for el in children[i_howto:i_somm]:
        body.remove(el)
    # 2) supprime le squelette (après le TOC jusqu'au sectPr)
    for el in children[i_toc + 1:]:
        if el is not sectpr:
            body.remove(el)
    # 3) nettoie les paragraphes vides en fin de page de garde (évite une page blanche)
    cover = [el for el in list(body)[: list(body).index(sommaire)] if el.tag == qn("w:p")]
    for el in reversed(cover):
        if p_text(el).strip():
            break
        body.remove(el)

    # saut de page avant le Sommaire pour le détacher proprement de la garde
    _set_page_break_before(sommaire)


def replace_toc_field(doc) -> bool:
    """Remplace le champ TOC encapsulé dans un `w:sdt` par un champ TOC standard.

    Le modèle stocke le sommaire dans un `w:sdt` (content control) au résultat vide :
    ni Word ni LibreOffice ne le reconstruisent de façon fiable. Un champ TOC « propre »
    (un seul paragraphe) est reconnu comme index natif → reconstruit par Word (à
    l'ouverture / F9) et par LibreOffice (cf. `bake_fields_libreoffice`).
    """
    body = doc.element.body
    for el in list(body):
        if el.tag == qn("w:sdt") and any("TOC" in (it.text or "") for it in el.iter(qn("w:instrText"))):
            xml = (
                f'<w:p {nsdecls("w")}>'
                '<w:r><w:fldChar w:fldCharType="begin" w:dirty="true"/></w:r>'
                '<w:r><w:instrText xml:space="preserve"> TOC \\o "1-3" \\h \\z \\u </w:instrText></w:r>'
                '<w:r><w:fldChar w:fldCharType="separate"/></w:r>'
                '<w:r><w:t xml:space="preserve">Sommaire — à mettre à jour (F9 / Outils ▸ Actualiser)</w:t></w:r>'
                '<w:r><w:fldChar w:fldCharType="end"/></w:r>'
                '</w:p>'
            )
            el.addprevious(parse_xml(xml))
            body.remove(el)
            return True
    return False


def mark_toc_dirty(doc) -> None:
    for el in doc.element.body.iter(qn("w:fldChar")):
        if el.get(qn("w:fldCharType")) == "begin":
            el.set(qn("w:dirty"), "true")
    # settings/updateFields = true (Word propose la mise à jour des champs à l'ouverture)
    settings = doc.settings.element
    if settings.find(qn("w:updateFields")) is None:
        uf = OxmlElement("w:updateFields")
        uf.set(qn("w:val"), "true")
        settings.append(uf)


def _find_soffice():
    for p in ("/Applications/LibreOffice.app/Contents/MacOS/soffice",
              "/usr/bin/soffice", "/usr/bin/libreoffice"):
        if os.path.isfile(p):
            return p
    for name in ("soffice", "libreoffice"):
        x = shutil.which(name)
        if x:
            return x
    return None


def bake_fields_libreoffice(docx_path, soffice=None, timeout=180) -> bool:
    """Reconstruit (« bake ») le sommaire et les champs via LibreOffice **en local**.

    Ouvre le `.docx` produit, met à jour les index (sommaire) et les champs, puis ré-enregistre :
    le sommaire est alors **peuplé dans le fichier** → visible et à jour dans **Word ET
    LibreOffice**, sans action de l'utilisateur. Aucun service externe. Repli sans perte si
    LibreOffice est absent (le champ TOC reste mis à jour par Word à l'ouverture / F9).
    """
    soffice = soffice or _find_soffice()
    if not soffice:
        sys.stderr.write("[toc] LibreOffice introuvable → sommaire mis à jour à l'ouverture (Word/F9).\n")
        return False
    path = os.path.abspath(docx_path).replace('"', '')
    prof = tempfile.mkdtemp(prefix="dat_lo_")
    std = os.path.join(prof, "user", "basic", "Standard")
    os.makedirs(std, exist_ok=True)
    module = (
        '<?xml version="1.0" encoding="UTF-8"?>\n'
        '<!DOCTYPE script:module PUBLIC "-//OpenOffice.org//DTD OfficeDocument 1.0//EN" "module.dtd">\n'
        '<script:module xmlns:script="http://openoffice.org/2000/script" script:name="Module1" script:language="StarBasic">\n'
        'Sub BakeToc\n'
        '  Dim oArgs(0) As New com.sun.star.beans.PropertyValue\n'
        '  oArgs(0).Name = "Hidden" : oArgs(0).Value = True\n'
        '  Dim oDoc\n'
        f'  oDoc = StarDesktop.loadComponentFromURL(ConvertToURL("{path}"), "_blank", 0, oArgs())\n'
        '  Dim oIdx, i\n'
        '  oIdx = oDoc.getDocumentIndexes()\n'
        '  For i = 0 To oIdx.Count - 1\n'
        '    oIdx.getByIndex(i).update()\n'
        '  Next i\n'
        '  oDoc.getTextFields().refresh()\n'
        '  oDoc.store()\n'
        '  oDoc.close(False)\n'
        '  StarDesktop.terminate()\n'
        'End Sub\n'
        '</script:module>\n'
    )
    with open(os.path.join(std, "Module1.xba"), "w", encoding="utf-8") as f:
        f.write(module)
    with open(os.path.join(std, "script.xlb"), "w", encoding="utf-8") as f:
        f.write('<?xml version="1.0" encoding="UTF-8"?>\n'
                '<!DOCTYPE library:library PUBLIC "-//OpenOffice.org//DTD OfficeDocument 1.0//EN" "library.dtd">\n'
                '<library:library xmlns:library="http://openoffice.org/2000/library" library:name="Standard" '
                'library:readonly="false" library:passwordprotected="false">\n'
                ' <library:element library:name="Module1"/>\n</library:library>\n')
    with open(os.path.join(prof, "user", "basic", "script.xlc"), "w", encoding="utf-8") as f:
        f.write('<?xml version="1.0" encoding="UTF-8"?>\n'
                '<!DOCTYPE library:libraries PUBLIC "-//OpenOffice.org//DTD OfficeDocument 1.0//EN" "libraries.dtd">\n'
                '<library:libraries xmlns:library="http://openoffice.org/2000/library" '
                'xmlns:xlink="http://www.w3.org/1999/xlink">\n'
                ' <library:library library:name="Standard" xlink:href="$(USER)/basic/Standard/script.xlb" '
                'xlink:type="simple" library:link="false"/>\n</library:libraries>\n')
    try:
        r = subprocess.run(
            [soffice, f"-env:UserInstallation=file://{prof}", "--headless", "--invisible",
             "--norestore", "--nologo",
             "vnd.sun.star.script:Standard.Module1.BakeToc?language=Basic&location=application"],
            capture_output=True, text=True, timeout=timeout)
        if r.returncode != 0:
            sys.stderr.write(f"[toc] bake LibreOffice rc={r.returncode} → repli (MAJ à l'ouverture).\n")
            return False
        # vérifie que le sommaire est réellement peuplé (sinon repli honnête)
        try:
            import zipfile as _zip
            doc = _zip.ZipFile(path).read("word/document.xml").decode("utf-8", "ignore")
            if "<w:hyperlink" not in doc:
                sys.stderr.write("[toc] bake sans effet (headless sans mise en page) → MAJ à l'ouverture.\n")
                return False
        except Exception:  # noqa: BLE001
            return False
        return True
    except Exception as e:  # noqa: BLE001
        sys.stderr.write(f"[toc] bake LibreOffice indisponible ({e}) → repli (MAJ à l'ouverture).\n")
        return False
    finally:
        shutil.rmtree(prof, ignore_errors=True)


def _fill_paragraphs(paragraphs, service, version, date) -> None:
    repl = {}
    if service:
        repl["[NOM DU SERVICE]"] = service
    if date:
        repl["[JJ mois AAAA]"] = date
    for para in paragraphs:
        full = "".join(r.text for r in para.runs)
        if not full.strip():
            continue
        new = full
        for k, v in repl.items():
            new = new.replace(k, v)
        # retire la mention « — Modèle / MODÈLE » (toutes casses, variantes de tiret)
        new = re.sub(r"\s*[—–-]\s*mod[èe]le\b", "", new, flags=re.IGNORECASE)
        if version and new.strip().startswith("Version"):
            new = re.sub(r"Version\s*:.*", f"Version : {version}", new)
        if new != full and para.runs:
            for r in para.runs[1:]:
                r.text = ""
            para.runs[0].text = new


def fill_cover(doc, service, version, date) -> None:
    """Renseigne garde + en-têtes/pieds (service, version, date) et retire « — Modèle ».

    Traite le corps ET les en-têtes/pieds de toutes les sections (le bandeau du
    modèle « DAT MirAI — Modèle » / « [NOM DU SERVICE] » vit dans l'en-tête).
    """
    _fill_paragraphs(doc.paragraphs, service, version, date)
    for section in doc.sections:
        for hf in (section.header, section.first_page_header, section.even_page_header,
                   section.footer, section.first_page_footer, section.even_page_footer):
            _fill_paragraphs(hf.paragraphs, service, version, date)


def _meta_version_rows(md):
    """Extrait les lignes de la table « Suivi des mises à jour » du bloc de méta.

    Le bloc de méta (avant le premier `## `) est sinon ignoré au rendu : on en
    récupère ici la table dont l'en-tête commence par « Version » (en-tête exclu,
    le modèle ayant déjà le sien).
    """
    data = []
    for line in md.splitlines():
        if line.startswith("## "):
            break
        st = line.strip()
        if not st.startswith("|"):
            continue
        if set(st) <= set("|-: "):  # ligne de séparation |---|---|
            continue
        data.append([c.strip() for c in st.strip("|").split("|")])
    if data and data[0] and data[0][0].lower().startswith("version"):
        data = data[1:]  # retire l'en-tête (déjà présent dans la table du modèle)
    return data


def fill_version_table(doc, md) -> bool:
    """Peuple la table « Suivi des mises à jour » de la garde avec celle du Markdown.

    Remplace les lignes de données (placeholders du modèle) par celles du Markdown,
    en conservant la ligne d'en-tête du modèle. Met le bloc « Suivi des mises à jour »
    sur une **page séparée** (saut de page avant le titre) et **réduit la police** de la
    table (`VERSION_TABLE_PT`). À appeler AVANT `render_markdown` (la table de garde est
    alors la première — et seule — table du document).
    """
    rows = _meta_version_rows(md)
    if not rows or not doc.tables:
        return False
    table = doc.tables[0]
    for row in list(table.rows)[1:]:        # garde l'en-tête, retire le reste
        row._tr.getparent().remove(row._tr)
    for values in rows:
        cells = table.add_row().cells
        for j, val in enumerate(values):
            if j < len(cells):
                _add_runs(cells[j].paragraphs[0], val)
    # police réduite + interligne compact sur TOUTE la table (en-tête + données)
    for trow in table.rows:
        for cell in trow.cells:
            for para in cell.paragraphs:
                para.paragraph_format.space_after = Pt(0)
                para.paragraph_format.line_spacing = 1.0
                for r in para.runs:
                    r.font.size = Pt(VERSION_TABLE_PT)
    # bloc « Suivi des mises à jour » sur une page séparée (saut de page avant le titre)
    for para in doc.paragraphs:
        if para.text.strip().lower().startswith("suivi des mises à jour"):
            _set_page_break_before(para._p)
            break
    return True


# ── Mise en page : polices réduites + aération ────────────────────────────────

def _find_style_el(doc, style_id):
    for st in doc.styles.element.findall(qn("w:style")):
        if st.get(qn("w:styleId")) == style_id:
            return st
    return None


def _ensure_child(parent, tag, before_tags=()):
    """Retourne l'enfant `tag`, le créant en respectant l'ordre du schéma si besoin."""
    el = parent.find(qn(tag))
    if el is not None:
        return el
    el = OxmlElement(tag)
    # insère avant le premier des before_tags présents (ordre OOXML), sinon append
    for bt in before_tags:
        sib = parent.find(qn(bt))
        if sib is not None:
            sib.addprevious(el)
            return el
    parent.append(el)
    return el


def _set_rpr_size(rpr, pt) -> None:
    for tag in ("w:sz", "w:szCs"):
        e = _ensure_child(rpr, tag)
        e.set(qn("w:val"), str(int(round(pt * 2))))  # demi-points


def _set_spacing(ppr, line=None, after=None, before=None) -> None:
    sp = _ensure_child(ppr, "w:spacing")
    if line is not None:
        sp.set(qn("w:line"), str(int(round(line * 240))))  # 240 = simple
        sp.set(qn("w:lineRule"), "auto")
    if after is not None:
        sp.set(qn("w:after"), str(int(round(after * 20))))   # points -> twips
    if before is not None:
        sp.set(qn("w:before"), str(int(round(before * 20))))


def apply_layout(doc) -> None:
    """Réduit les polices et aère le document, en surchargeant les styles du modèle.

    Agit sur `docDefaults` (corps), `Normal` si présent, et `Heading1/2/3`.
    """
    styles = doc.styles.element
    dd = styles.find(qn("w:docDefaults"))
    if dd is not None:
        rprd = _ensure_child(dd, "w:rPrDefault")
        _set_rpr_size(_ensure_child(rprd, "w:rPr"), BODY_PT)
        pprd = _ensure_child(dd, "w:pPrDefault")
        ppr = _ensure_child(pprd, "w:pPr")
        _set_spacing(ppr, line=LINE_SPACING, after=PARA_AFTER_PT)
        # contrôle veuves/orphelines : pas de ligne seule en haut/bas de page
        wc = _ensure_child(ppr, "w:widowControl")
        wc.set(qn("w:val"), "true")
    normal = _find_style_el(doc, "Normal")
    if normal is not None:
        _set_rpr_size(_ensure_child(normal, "w:rPr", before_tags=()), BODY_PT)
        _set_spacing(_ensure_child(normal, "w:pPr", before_tags=("w:rPr",)),
                     line=LINE_SPACING, after=PARA_AFTER_PT)
    for sid, pt in (("Heading1", H1_PT), ("Heading2", H2_PT), ("Heading3", H3_PT)):
        st = _find_style_el(doc, sid)
        if st is None:
            continue
        _set_rpr_size(_ensure_child(st, "w:rPr", before_tags=()), pt)
        before, after = HEADING_SPACING.get(sid, (None, None))
        _set_spacing(_ensure_child(st, "w:pPr", before_tags=("w:rPr",)), before=before, after=after)


# ── Diagrammes mermaid → image (rendu LOCAL, aucun service externe) ────────────

def _find_chrome():
    for p in CHROME_CANDIDATES:
        if p and os.path.isfile(p):
            return p
    for name in ("google-chrome", "chromium", "chromium-browser", "chrome"):
        p = shutil.which(name)
        if p:
            return p
    return None


def _mermaid_cmd():
    if shutil.which("mmdc"):
        return ["mmdc"]
    if shutil.which("npx"):
        return ["npx", "-y", "@mermaid-js/mermaid-cli"]
    return None


def render_mermaid_png(src, out_png, chrome=None, npm_cache=None) -> bool:
    """Rend un diagramme mermaid en PNG via mermaid-cli + Chrome LOCAL. False si indispo."""
    cmd = _mermaid_cmd()
    if not cmd:
        return False
    tmp_mmd = tmp_cfg = None
    try:
        with tempfile.NamedTemporaryFile("w", suffix=".mmd", delete=False, encoding="utf-8") as f:
            f.write(src)
            tmp_mmd = f.name
        cfg = {
            "theme": "neutral",
            "themeVariables": {"fontFamily": "Arial", "fontSize": "15px"},
            "flowchart": {"useMaxWidth": True, "htmlLabels": True},
            "sequence": {"useMaxWidth": True},
        }
        with tempfile.NamedTemporaryFile("w", suffix=".json", delete=False, encoding="utf-8") as cf:
            json.dump(cfg, cf)
            tmp_cfg = cf.name
        env = dict(os.environ)
        if chrome:
            env["PUPPETEER_EXECUTABLE_PATH"] = chrome
        if npm_cache:
            env["npm_config_cache"] = npm_cache
        args = cmd + ["-i", tmp_mmd, "-o", out_png, "-c", tmp_cfg, "-s", "2", "-b", "white"]
        r = subprocess.run(args, env=env, capture_output=True, text=True, timeout=300)
        if r.returncode != 0 or not os.path.isfile(out_png):
            sys.stderr.write(f"[mermaid] échec (rc={r.returncode}) : {(r.stderr or '')[-400:]}\n")
            return False
        return True
    except Exception as e:  # noqa: BLE001 — repli silencieux vers le rendu texte
        sys.stderr.write(f"[mermaid] exception : {e}\n")
        return False
    finally:
        for t in (tmp_mmd, tmp_cfg):
            if t:
                try:
                    os.unlink(t)
                except OSError:
                    pass


def _content_width(doc):
    sect = doc.sections[0]
    return sect.page_width - sect.left_margin - sect.right_margin


def _insert_image(doc, png) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pic = p.add_run().add_picture(png, width=_content_width(doc))
    max_h = Inches(MERMAID_MAX_H_IN)
    if pic.height > max_h:
        ratio = max_h / pic.height
        pic.height = int(pic.height * ratio)
        pic.width = int(pic.width * ratio)


def insert_cover_logo(doc, logo_path, width_cm=LOGO_WIDTH_CM) -> bool:
    """Insère le logo (PNG) centré en tête de page de garde. False si asset absent.

    Le logo devient le tout premier élément du corps (avant « MINISTÈRE DE
    L'INTÉRIEUR — DTNUM », issu du modèle). Repli sans perte si le fichier manque.
    """
    if not logo_path or not os.path.isfile(logo_path):
        sys.stderr.write(f"[logo] introuvable ({logo_path}) → page de garde sans logo.\n")
        return False
    p = doc.add_paragraph()  # ajouté en fin de corps…
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(12)
    p.add_run().add_picture(logo_path, width=Cm(width_cm))
    doc.element.body.insert(0, p._p)  # …puis déplacé en tête du corps
    return True


# ── Styles & mise en forme ────────────────────────────────────────────────────

def _set_style_id(paragraph, style_id) -> None:
    """Applique un style par son styleId réel (robuste vs accès par nom)."""
    pPr = paragraph._p.get_or_add_pPr()
    for old in pPr.findall(qn("w:pStyle")):
        pPr.remove(old)
    pStyle = OxmlElement("w:pStyle")
    pStyle.set(qn("w:val"), style_id)
    pPr.insert(0, pStyle)


def _set_page_break_before(el) -> None:
    pPr = el.find(qn("w:pPr"))
    if pPr is None:
        pPr = OxmlElement("w:pPr")
        el.insert(0, pPr)
    if pPr.find(qn("w:pageBreakBefore")) is None:
        pPr.append(OxmlElement("w:pageBreakBefore"))


# ── Typographie française (Imprimerie nationale) — idempotente, 100 % locale ──

_TYPO_PROTECT_RE = re.compile(
    r"`[^`]*`"                                # code inline `…`
    r"|(?:https?://|ftp://|mailto:)\S+"       # URLs
    r"|\\\\[^\s]+"                             # chemins UNC \\serveur\…
    r"|[\w.+-]+@[\w.-]+\.\w+"                  # e-mails
)


def _french_double_quotes(s: str) -> str:
    """Paires équilibrées de `"` droits → « … » (fines insécables intérieures).

    Un `"` orphelin (nombre impair) est laissé tel quel.
    """
    if s.count('"') < 2:
        return s
    pairable = s.count('"')
    if pairable % 2:
        pairable -= 1  # laisse le dernier guillemet impair intact
    out, open_next, used = [], True, 0
    for ch in s:
        if ch == '"' and used < pairable:
            out.append("«" + NNBSP if open_next else NNBSP + "»")
            open_next = not open_next
            used += 1
        else:
            out.append(ch)
    return "".join(out)


def normalize_typography(text: str) -> str:
    """Applique la typographie française à de la prose. **Idempotent.**

    Règle d'or : on ne CONVERTIT qu'une espace déjà présente, jamais on n'en ajoute
    (préserve `12:30`, `clé:valeur`, URLs, ratios…). Conventions Imprimerie nationale :
    NBSP avant `:`, fine insécable avant `;` `!` `?` et autour des guillemets français.
    Les segments code inline / URL / e-mail / UNC sont protégés (placeholders).
    """
    if not text:
        return text
    # 0) Protéger les segments à ne pas typographier
    protected = []

    def _stash(m):
        protected.append(m.group(0))
        return f"\x00{len(protected) - 1}\x00"

    s = _TYPO_PROTECT_RE.sub(_stash, text)

    # 1) Apostrophe droite → typographique (sans conflit FR/EN)
    s = s.replace("'", "’")
    # 2) Guillemets doubles droits → « … »
    s = _french_double_quotes(s)
    # 3) Séparateur de milliers : 30 000 → fine insécable
    s = re.sub(r"\d{1,3}(?: \d{3})+", lambda m: m.group(0).replace(" ", NNBSP), s)
    # 4) Nombre + unité → insécable (espace existante uniquement)
    units = "|".join(sorted((re.escape(u) for u in TYPO_UNITS), key=len, reverse=True))
    s = re.sub(rf"(\d) ({units})\b", rf"\1{NBSP}\2", s)
    # 5) Ponctuation haute (convertit l'espace régulière existante)
    s = re.sub(r" :", NBSP + ":", s)            # NBSP avant deux-points
    s = re.sub(r" ([;!?])", NNBSP + r"\1", s)   # fine avant ; ! ?
    s = re.sub(r" ([%€°‰])", NBSP + r"\1", s)   # insécable avant % € ° ‰
    # 6) Guillemets français déjà présents : fine insécable intérieure
    s = s.replace("« ", "«" + NNBSP).replace(" »", NNBSP + "»")
    # 7) Repli éventuel des fines sur insécables normales (police « tofu »)
    if THIN_NBSP_FALLBACK:
        s = s.replace(NNBSP, NBSP)
    # 8) Restaurer les segments protégés
    s = re.sub(r"\x00(\d+)\x00", lambda m: protected[int(m.group(1))], s)
    return s


def _add_runs(paragraph, text) -> None:
    """Rend le code inline (chasse fixe), le gras `**…**`, l'italique `*…*` et la typographie.

    L'**emphase peut envelopper du code inline** (ex. `*(source : `chemin`)*`) : le code est
    d'abord remplacé par des marqueurs, la typographie est appliquée à la prose (le code reste
    intact), puis l'emphase est analysée — chaque segment ré-injecte le code en chasse fixe.
    Les astérisques échappées `\\*` sont rendues littéralement.
    """
    codes = []

    def _stash_code(m):
        codes.append(m.group(1))
        return f"\x02{len(codes) - 1}\x02"

    s = re.sub(r"`([^`]*)`", _stash_code, text)   # 1) protéger le code inline
    s = normalize_typography(s)                    # 2) typographie (code protégé)
    s = s.replace("\\*", "\x01")                   # 3) protéger \* (astérisque littérale)

    def emit(chunk, bold=False, italic=False):
        # ré-injecte les marqueurs de code en runs à chasse fixe
        for j, part in enumerate(re.split(r"\x02(\d+)\x02", chunk)):
            mono = j % 2 == 1
            txt = codes[int(part)] if mono else part.replace("\x01", "*")
            if not txt:
                continue
            run = paragraph.add_run(txt)
            if mono:
                run.font.name = "Consolas"
            if bold:
                run.bold = True
            if italic:
                run.italic = True

    pos = 0
    for m in re.finditer(r"\*\*(.+?)\*\*|\*(.+?)\*", s):
        if m.start() > pos:
            emit(s[pos:m.start()])
        if m.group(1) is not None:
            emit(m.group(1), bold=True)
        else:
            emit(m.group(2), italic=True)
        pos = m.end()
    if pos < len(s):
        emit(s[pos:])


def _set_table_borders(table) -> None:
    tblPr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{edge}")
        e.set(qn("w:val"), "single")
        e.set(qn("w:sz"), "4")
        e.set(qn("w:space"), "0")
        e.set(qn("w:color"), "999999")
        borders.append(e)
    tblPr.append(borders)


def _add_table(doc, rows):
    header, body_rows = rows[0], rows[1:]
    table = doc.add_table(rows=1, cols=len(header))
    try:
        table.style = "Table Grid"
    except Exception:
        _set_table_borders(table)
    for j, cell in enumerate(header):
        para = table.rows[0].cells[j].paragraphs[0]
        _add_runs(para, cell)
        for r in para.runs:
            r.bold = True
    # en-tête répété + ligne non sécable
    trPr = table.rows[0]._tr.get_or_add_trPr()
    for tag in ("w:tblHeader", "w:cantSplit"):
        trPr.append(OxmlElement(tag))
    for row in body_rows:
        cells = table.add_row().cells
        tr = table.rows[-1]._tr
        tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
        for j, val in enumerate(row):
            if j < len(cells):
                _add_runs(cells[j].paragraphs[0], val)
    # tableaux compacts : police réduite + interligne simple, pas d'espace après
    for trow in table.rows:
        for cell in trow.cells:
            for para in cell.paragraphs:
                para.paragraph_format.space_after = Pt(0)
                para.paragraph_format.line_spacing = 1.0
                for r in para.runs:
                    r.font.size = Pt(TABLE_PT)
    return table


# ── Parsing Markdown → corps Word ─────────────────────────────────────────────

def _split_table_row(line):
    cells = [c.strip() for c in line.strip().strip("|").split("|")]
    return cells


def render_markdown(doc, md, opts=None) -> None:
    if opts is None:
        opts = SimpleNamespace(mermaid=False, tmp=None, chrome=None, npm_cache=None, counter=[0])
    lines = md.splitlines()
    started = False  # on commence au premier '## '
    i, n = 0, len(lines)
    while i < n:
        line = lines[i]
        h = re.match(r"^(#{1,6})\s+(.*)$", line)
        if h and len(h.group(1)) >= 2:
            started = True
        if not started:
            i += 1
            continue

        # Titres de section
        if h:
            level = len(h.group(1))
            style = LEVEL_STYLE.get(level)
            if style:
                p = doc.add_paragraph()
                _set_style_id(p, style)
                _add_runs(p, h.group(2).strip())
                p.paragraph_format.keep_with_next = True
                if style == "Heading1":
                    _set_page_break_before(p._p)
            i += 1
            continue

        # Bloc mermaid ```mermaid -> image (rendu local) ; repli : chasse fixe
        if line.strip().startswith("```mermaid"):
            i += 1
            src_lines = []
            while i < n and not lines[i].strip().startswith("```"):
                src_lines.append(lines[i])
                i += 1
            i += 1  # saute la fence de fermeture
            png = None
            if opts.mermaid and opts.tmp:
                png = os.path.join(opts.tmp, f"mermaid_{opts.counter[0]}.png")
                opts.counter[0] += 1
                if not render_mermaid_png("\n".join(src_lines), png, opts.chrome, opts.npm_cache):
                    png = None
            if png:
                _insert_image(doc, png)
            else:  # repli : source mermaid en chasse fixe (rien n'est perdu)
                for src_line in src_lines:
                    p = doc.add_paragraph()
                    run = p.add_run(src_line)
                    run.font.name = "Consolas"
                    run.font.size = Pt(CODE_PT)
            continue

        # Bloc de code ``` -> police à chasse fixe
        if line.strip().startswith("```"):
            i += 1
            while i < n and not lines[i].strip().startswith("```"):
                p = doc.add_paragraph()
                run = p.add_run(lines[i])
                run.font.name = "Consolas"
                run.font.size = Pt(CODE_PT)
                i += 1
            i += 1
            continue

        # Tableau Markdown
        if line.lstrip().startswith("|") and i + 1 < n and re.match(r"^\s*\|?[\s:|-]+\|", lines[i + 1]):
            rows = [_split_table_row(line)]
            i += 2  # saute le séparateur
            while i < n and lines[i].lstrip().startswith("|"):
                rows.append(_split_table_row(lines[i]))
                i += 1
            _add_table(doc, rows)
            doc.add_paragraph()
            continue

        # Liste à puces : trait d'union + retrait pendant (lisibilité)
        mli = re.match(r"^\s*[-*]\s+(.*)$", line)
        if mli:
            p = doc.add_paragraph()
            _set_style_id(p, "ListParagraph")
            pf = p.paragraph_format
            pf.left_indent = Cm(LIST_INDENT_CM)
            pf.first_line_indent = Cm(-LIST_HANGING_CM)  # retrait pendant
            pf.space_after = Pt(LIST_AFTER_PT)
            _add_runs(p, BULLET + mli.group(1))
            i += 1
            continue

        # Citation / encadré
        mq = re.match(r"^>\s?(.*)$", line)
        if mq:
            p = doc.add_paragraph()
            _add_runs(p, mq.group(1))
            for r in p.runs:
                r.italic = True
            i += 1
            continue

        # Règle horizontale -> ignorée
        if re.match(r"^\s*---+\s*$", line):
            i += 1
            continue

        # Paragraphe normal
        if line.strip():
            p = doc.add_paragraph()
            _add_runs(p, line.strip())
        i += 1


def reorder_sectpr(doc) -> None:
    body = doc.element.body
    sectpr = body.find(qn("w:sectPr"))
    if sectpr is not None:
        body.remove(sectpr)
        body.append(sectpr)


# ── Main ──────────────────────────────────────────────────────────────────────

def main() -> None:
    ap = argparse.ArgumentParser(description="Rend un DAT MirAI en .docx sur le modèle ministériel.")
    ap.add_argument("--input", required=True, help="Contenu Markdown fidèle à la structure du modèle")
    ap.add_argument("--output", required=True, help="Chemin du .docx de sortie")
    here = os.path.dirname(os.path.abspath(__file__))
    default_model = os.path.normpath(os.path.join(here, "..", "dat-generation", "dat", "MODELE_DAT_MirAI.docx"))
    ap.add_argument("--model", default=default_model, help="Modèle .docx (défaut : skill dat-generation)")
    ap.add_argument("--service", default="", help="Nom du service (page de garde)")
    ap.add_argument("--version", default="", help="Version (page de garde)")
    ap.add_argument("--date", default="", help="Date (page de garde)")
    ap.add_argument("--chrome", default=None,
                    help="Chemin du binaire Chrome/Chromium pour le rendu mermaid (sinon auto-détection)")
    ap.add_argument("--npm-cache", default=None,
                    help="Cache npm pour `npx mermaid-cli` (utile en environnement sandboxé)")
    ap.add_argument("--no-mermaid", action="store_true",
                    help="Désactive le rendu image mermaid (repli texte à chasse fixe)")
    ap.add_argument("--logo", default=LOGO_DEFAULT,
                    help="Logo PNG de la page de garde (défaut : assets/logo_minint_2020.png)")
    ap.add_argument("--no-logo", action="store_true",
                    help="Ne pas insérer de logo en page de garde")
    ap.add_argument("--soffice", default=None,
                    help="Chemin LibreOffice (soffice) pour --toc-bake (sinon auto-détection)")
    ap.add_argument("--toc-bake", action="store_true",
                    help="Tente de pré-remplir le sommaire via LibreOffice (nécessite un LibreOffice "
                         "capable de mise en page ; peu fiable en headless). Sinon le champ TOC est "
                         "mis à jour à l'ouverture (Word) ou via Outils ▸ Actualiser (LibreOffice).")
    args = ap.parse_args()

    if not os.path.isfile(args.model):
        sys.exit(f"Modèle introuvable : {args.model}")
    with open(args.input, encoding="utf-8") as f:
        md = f.read()

    mermaid_on = not args.no_mermaid
    chrome = args.chrome or _find_chrome()
    if mermaid_on and _mermaid_cmd() is None:
        sys.stderr.write("[mermaid] mmdc/npx introuvable → repli texte pour les diagrammes.\n")
        mermaid_on = False
    elif mermaid_on and chrome is None:
        sys.stderr.write("[mermaid] Chrome/Chromium introuvable → repli texte pour les diagrammes.\n")
        mermaid_on = False

    doc = Document(args.model)
    apply_layout(doc)
    strip_template_body(doc)
    fill_cover(doc, args.service, args.version, args.date)
    fill_version_table(doc, md)  # garde : table « Suivi des mises à jour » depuis le Markdown
    replace_toc_field(doc)       # sommaire : champ TOC standard (reconnu Word + LibreOffice)
    if not args.no_logo:
        insert_cover_logo(doc, args.logo)

    tmp = tempfile.mkdtemp(prefix="dat_mermaid_") if mermaid_on else None
    try:
        opts = SimpleNamespace(mermaid=mermaid_on, tmp=tmp, chrome=chrome,
                               npm_cache=args.npm_cache, counter=[0])
        render_markdown(doc, md, opts)
        reorder_sectpr(doc)
        mark_toc_dirty(doc)
        os.makedirs(os.path.dirname(os.path.abspath(args.output)), exist_ok=True)
        doc.save(args.output)
    finally:
        if tmp:
            shutil.rmtree(tmp, ignore_errors=True)
    # Sommaire : champ TOC standard (reconnu Word + LibreOffice), mis à jour à l'ouverture.
    # Option --toc-bake : tente un pré-remplissage via LibreOffice (vérifié, sinon repli).
    baked = bake_fields_libreoffice(args.output, args.soffice) if args.toc_bake else False
    print(f"OK → {args.output}"
          + (f"  ({opts.counter[0]} diagramme(s) mermaid)" if mermaid_on else "")
          + ("  + sommaire pré-rempli" if baked else "  (sommaire : MAJ à l'ouverture / Outils ▸ Actualiser)"))


if __name__ == "__main__":
    main()
